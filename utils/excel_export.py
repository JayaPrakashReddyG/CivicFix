
from pathlib import Path
from datetime import datetime

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment


BASE_DIR = Path(__file__).resolve().parent.parent
EXCEL_FILE = BASE_DIR / "CivicFix_Registered_Users.xlsx"


def _get_workbook():
    """Existing workbook open cheyyi, lekapothe create cheyyi."""
    if EXCEL_FILE.exists():
        workbook = load_workbook(EXCEL_FILE)
    else:
        workbook = Workbook()
        default_sheet = workbook.active
        workbook.remove(default_sheet)

    # Registered Users sheet
    if "Registered Users" not in workbook.sheetnames:
        users_sheet = workbook.create_sheet("Registered Users")
        users_sheet.append([
            "User ID",
            "Name",
            "Email",
            "Role",
            "Registered At",
        ])
        _format_sheet(users_sheet)

    # Login History sheet
    if "Login History" not in workbook.sheetnames:
        login_sheet = workbook.create_sheet("Login History")
        login_sheet.append([
            "User ID",
            "Name",
            "Email",
            "Login At",
        ])
        _format_sheet(login_sheet)

    return workbook


def _format_sheet(sheet):
    """Header styling and column widths."""
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(
            fill_type="solid",
            fgColor="1565C0",
        )
        cell.alignment = Alignment(horizontal="center")

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions

    for column, width in {
        "A": 12,
        "B": 28,
        "C": 36,
        "D": 18,
        "E": 24,
    }.items():
        sheet.column_dimensions[column].width = width


def export_registered_user(user):
    """Register ayina user details Excel lo add cheyyi."""
    if not user:
        raise ValueError("User details are missing.")

    user_id = user.get("id")
    name = user.get("name", "")
    email = user.get("email", "")
    role = user.get("role", "citizen")
    registered_at = (
        user.get("created_at")
        or datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    )

    if not user_id or not email:
        raise ValueError("User ID or email is missing.")

    workbook = _get_workbook()
    sheet = workbook["Registered Users"]

    existing_emails = {
        str(sheet.cell(row=row, column=3).value).strip().lower()
        for row in range(2, sheet.max_row + 1)
        if sheet.cell(row=row, column=3).value
    }

    if email.strip().lower() not in existing_emails:
        sheet.append([
            user_id,
            name,
            email,
            role,
            str(registered_at),
        ])

    _format_sheet(sheet)
    workbook.save(EXCEL_FILE)
    workbook.close()


def export_login_history(user):
    """Successful login details ni separate sheet lo add cheyyi."""
    if not user:
        raise ValueError("User details are missing.")

    user_id = user.get("user_id", user.get("id"))
    name = user.get("name", "")
    email = user.get("email", "")
    login_at = (
        user.get("login_at")
        or datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    )

    if not user_id or not email:
        raise ValueError("User ID or email is missing.")

    workbook = _get_workbook()
    sheet = workbook["Login History"]

    # Prathi successful login ki kottha row add avuthundi.
    sheet.append([
        user_id,
        name,
        email,
        str(login_at),
    ])

    _format_sheet(sheet)
    workbook.save(EXCEL_FILE)
    workbook.close()
